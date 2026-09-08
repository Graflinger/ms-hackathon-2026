import csv
import io
import json
import zipfile
from datetime import date, datetime
from pathlib import PurePosixPath

from goldenloop_eval import Case, Turn
from openpyxl import load_workbook

from .config import Settings, clean


class ImportError(ValueError):
    pass


def parse_upload(raw: bytes, filename: str, sheet: str | None, settings: Settings) -> dict:
    suffix = PurePosixPath(filename.lower()).suffix
    tables = {}
    if suffix == ".csv":
        try:
            text = raw.decode("utf-8-sig")
            if "\x00" in text:
                raise ImportError("NUL bytes are not permitted")
            reader = csv.reader(io.StringIO(text, newline=""), strict=True)
            tables["CSV"] = []
            for row in reader:
                if len(tables["CSV"]) >= settings.max_rows + 1:
                    raise ImportError("Row limit exceeded")
                tables["CSV"].append([(cell, cell.lstrip().startswith(("=", "+", "@"))) for cell in row])
        except (UnicodeError, csv.Error) as exc:
            raise ImportError("CSV must be well-formed UTF-8") from exc
    elif suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                names = [entry.filename.lower() for entry in entries]
                if len(entries) > 1000 or len(names) != len(set(names)):
                    raise ImportError("Invalid or overly complex workbook archive")
                if any(
                    "vbaproject" in name or "macrosheet" in name or "externallinks/" in name for name in names
                ):
                    raise ImportError("Macros and external links are not permitted")
                if sum(entry.file_size for entry in entries) > 20 * 1024 * 1024:
                    raise ImportError("Workbook expanded size limit exceeded")
                for entry in entries:
                    if (
                        entry.flag_bits & 1
                        or ".." in PurePosixPath(entry.filename).parts
                        or entry.filename.startswith(("/", "\\"))
                        or entry.file_size > max(1024 * 1024, entry.compress_size * 200)
                    ):
                        raise ImportError("Unsafe workbook archive")
                types = archive.read("[Content_Types].xml").lower()
                if b"macroenabled" in types or b"vbaproject" in types:
                    raise ImportError("Macro-enabled workbooks are not permitted")
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.sheetnames) > 10:
                    raise ImportError("Workbook sheet limit exceeded")
                total = 0
                for worksheet in workbook.worksheets:
                    if (worksheet.max_row or 0) > settings.max_rows + 1 or (worksheet.max_column or 0) > 100:
                        raise ImportError("Workbook dimensions exceed limits")
                    table = []
                    for row in worksheet.iter_rows():
                        total += 1
                        if total > settings.max_rows + 10:
                            raise ImportError("Workbook total row limit exceeded")
                        table.append([(cell.value, cell.data_type == "f") for cell in row])
                    tables[worksheet.title] = table
            finally:
                workbook.close()
        except ImportError:
            raise
        except Exception as exc:
            raise ImportError("Invalid XLSX workbook") from exc
    else:
        raise ImportError("Only .csv and .xlsx files are accepted; macros are prohibited")

    parsed = {}
    for name, table in tables.items():
        errors = []
        if not table:
            parsed[name] = {"columns": [], "rows": [], "errors": ["Empty sheet"]}
            continue
        columns = [str(value).strip() if value is not None else "" for value, _ in table[0]]
        if len(columns) > 100 or not all(columns) or len(set(columns)) != len(columns):
            errors.append("Headers must be nonempty, unique, and at most 100 columns")
        if any(formula for _, formula in table[0]):
            errors.append("Formula headers are not permitted")
        rows = []
        for index, cells in enumerate(table[1:], 2):
            if not any(value is not None and str(value).strip() for value, _ in cells):
                continue
            row_errors = []
            if len(cells) != len(columns):
                row_errors.append("Column count does not match headers")
            if any(formula for _, formula in cells):
                row_errors.append("Formula cells are not permitted and were not evaluated")
            values = {}
            for column, (value, formula) in zip(columns, cells):
                if isinstance(value, (date, datetime)):
                    value = value.isoformat()
                if value is not None and len(str(value)) > 16000:
                    row_errors.append("Cell exceeds 16000 characters")
                values[column] = "[FORMULA NOT EVALUATED]" if formula else value
            rows.append({"row": index, "values": values, "errors": row_errors})
        if not rows:
            errors.append("No nonempty data rows")
        parsed[name] = {"columns": columns, "rows": rows, "errors": errors}
    selected = sheet or next(iter(parsed), None)
    if selected not in parsed:
        raise ImportError("Unknown sheet")
    # Persist only sanitized previews, never raw uploads or executable cell content.
    payload = {
        "filename": PurePosixPath(filename.replace("\\", "/")).name,
        "selected": selected,
        "sheets": list(parsed),
        "tables": parsed,
    }
    sanitized = clean(payload)
    # Column/sheet names are mapping keys. Redaction of a key must not corrupt the stored preview shape.
    if sanitized["selected"] not in sanitized["tables"] or any(
        not isinstance(table, dict)
        or not isinstance(table.get("rows"), list)
        or any(set(row["values"]) != set(table["columns"]) for row in table["rows"])
        for table in sanitized["tables"].values()
    ):
        raise ImportError("Sensitive sheet or column names are not supported; rename them before importing")
    return sanitized


def preview(payload: dict, sheet: str | None = None) -> dict:
    selected = sheet or payload["selected"]
    if selected not in payload["tables"]:
        raise ImportError("Unknown sheet")
    return {**payload["tables"][selected], "sheets": payload["sheets"]}


def mapped_cases(payload: dict, mapping: dict, sheet: str | None, import_id: str) -> list[Case]:
    selected = sheet or payload["selected"]
    table = preview(payload, selected)
    allowed = {
        "id",
        "case_id",
        "title",
        "user",
        "question",
        "reference_answer",
        "context",
        "scenario",
        "tags",
        "scenario_id",
        "turn_index",
    }
    if not mapping or set(mapping) - allowed or set(mapping.values()) - set(table["columns"]):
        raise ImportError("Invalid mapping fields or uploaded columns")
    for a, b in (("user", "question"), ("id", "case_id"), ("context", "scenario")):
        if a in mapping and b in mapping:
            raise ImportError(f"Map only one of {a} and {b}")
    if not ({"user", "question"} & set(mapping)):
        raise ImportError("Map a question or user column")
    if ("scenario_id" in mapping) != ("turn_index" in mapping):
        raise ImportError("Multi-turn mapping requires both scenario_id and turn_index")
    if table["errors"] or any(row["errors"] for row in table["rows"]):
        raise ImportError("Resolve all preview errors before committing")
    groups = []
    seen = set()
    for row in table["rows"]:
        item = {key: row["values"].get(column) for key, column in mapping.items()}
        item = {key: "" if value is None else str(value).strip() for key, value in item.items()}
        user = item.get("user", item.get("question", ""))
        if not user:
            raise ImportError(f"Row {row['row']}: question is required")
        scenario_id = item.get("scenario_id")
        if "scenario_id" in mapping:
            if not scenario_id or not item["turn_index"].isdigit():
                raise ImportError(f"Row {row['row']}: scenario ID and integer turn index required")
            turn = int(item["turn_index"])
            if not groups or groups[-1]["scenario_id"] != scenario_id:
                if scenario_id in seen:
                    raise ImportError("Scenario rows must be contiguous")
                if turn != 0:
                    raise ImportError("Scenario turn indexes must begin at zero")
                seen.add(scenario_id)
                groups.append({"scenario_id": scenario_id, "items": []})
            if turn != len(groups[-1]["items"]):
                raise ImportError("Scenario turn indexes must be contiguous and zero-based")
        else:
            groups.append({"scenario_id": None, "items": []})
        groups[-1]["items"].append((item, row["row"], user))
    cases = []
    for group in groups:
        first, _, user = group["items"][0]
        for item, row, _ in group["items"][1:]:
            for key in ("id", "case_id", "title", "context", "scenario", "tags"):
                if item.get(key) and item.get(key) != first.get(key):
                    raise ImportError(f"Row {row}: conflicting scenario metadata")
        tags = first.get("tags", "")
        if tags.startswith("["):
            try:
                tags = json.loads(tags)
                if not isinstance(tags, list) or not all(isinstance(tag, str) for tag in tags):
                    raise ValueError
            except (ValueError, TypeError) as exc:
                raise ImportError("Tags must be a string array or comma-separated text") from exc
        else:
            tags = [tag.strip() for tag in tags.split(",") if tag.strip()]
        case_id = first.get("id", first.get("case_id"))
        cases.append(
            Case(
                **({"id": case_id} if case_id else {}),
                title=first.get("title") or user[:200],
                tags=list(dict.fromkeys(["synthetic", *tags])),
                context=first.get("context", first.get("scenario", "")),
                source={
                    "type": "import",
                    "import_id": import_id,
                    "filename": payload["filename"],
                    "sheet": selected,
                    "rows": [row for _, row, _ in group["items"]],
                    "scenario_id": group["scenario_id"],
                    "synthetic": True,
                },
                turns=[
                    Turn(user=text, reference_answer=item.get("reference_answer") or None)
                    for item, _, text in group["items"]
                ],
            )
        )
    return cases
